"""Judges round trips of workbooks with openpyxl (MIT): what each copy kept of its original.

  python judge.py <originals folder> <copies folder> <out.json>

For every copy: whether it opens; every cell's value and formula; number formats, fonts,
fills, borders, alignment, merged cells, column widths, row heights, sheet names and order;
and the things an editor must leave intact: charts, pivot tables, conditional formats, data
validation, defined names, plus tables, comments, pictures, sparklines and slicers. A copy
made after an edit (an <name>.json beside it) leaves the edited cell and its dependents'
values out of the value check, and only the edited cell's formula may differ.
"""
import json, os, re, subprocess, sys, zipfile, warnings, hashlib
from concurrent.futures import ThreadPoolExecutor
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

def part_hashes(path, pattern):
    z = zipfile.ZipFile(path)
    return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if re.match(pattern, n)}

def raw_count(path, regex):
    z = zipfile.ZipFile(path); total = 0
    for n in z.namelist():
        if n.startswith("xl/worksheets/") and n.endswith(".xml"):
            total += len(re.findall(regex, z.read(n)))
    return total

FEATURES = {
    "charts": r"xl/charts/chart\d+\.xml$",
    "pivots": r"xl/pivot(Tables/pivotTable|Cache/pivotCacheDefinition)\d+\.xml$",
    "tables": r"xl/tables/table\d+\.xml$",
    "comments": r"xl/(comments\d+|threadedComments/.*)\.xml$",
    "pictures": r"xl/media/.*",
    "drawings": r"xl/drawings/drawing\d+\.xml$",
    "slicers": r"xl/(slicers|slicerCaches|timelines|timelineCaches)/.*\.xml$",
    "vba": r"xl/vbaProject\.bin$",
    "externalLinks": r"xl/externalLinks/.*\.xml$",
}
RAW = {"x14 conditional formats": rb"<x14:conditionalFormatting\b", "x14 validations": rb"<x14:dataValidation\b",
       "sparklines": rb"<x14:sparklineGroup\b"}

def color(c):
    if c is None: return None
    return (c.type, c.rgb if c.type == "rgb" else c.theme if c.type == "theme" else c.indexed, round(c.tint or 0, 6))

def look(cell):
    f, fl, b, a = cell.font, cell.fill, cell.border, cell.alignment
    return (cell.number_format,
            (f.name, f.sz, f.b, f.i, f.u, f.strike, color(f.color)),
            (fl.fill_type, color(getattr(fl, "fgColor", None)), color(getattr(fl, "bgColor", None))) if fl is not None else None,
            tuple((s.style, color(s.color)) if s is not None else None for s in (b.left, b.right, b.top, b.bottom)),
            (a.horizontal, a.vertical, a.wrap_text, a.indent, a.text_rotation))

def formula_of(v):
    if isinstance(v, ArrayFormula): return ("array", v.ref, v.text)
    return v

def same_value(a, b):
    if isinstance(a, float) or isinstance(b, float):
        try: return abs(float(a) - float(b)) <= 1e-9 * max(1, abs(float(a)), abs(float(b)))
        except Exception: return False
    if a in (None, "") and b in (None, ""): return True
    return a == b

def cf_rules(ws):
    out = []
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            out.append((str(rng.sqref), rule.type, rule.operator, tuple(rule.formula or ()), rule.priority))
    return sorted(out, key=str)

def dv_rules(ws):
    return sorted(((str(d.sqref), d.type, d.formula1, d.formula2, d.operator) for d in ws.data_validations.dataValidation), key=str)

def names(wb):
    out = {}
    for n, d in wb.defined_names.items(): out[("", n)] = d.attr_text
    for ws in wb.worksheets:
        for n, d in getattr(ws, "defined_names", {}).items(): out[(ws.title, n)] = d.attr_text
    return out

def judge(job):
    original, copy, edit = job
    r = {"file": os.path.basename(copy)}
    try:
        o = openpyxl.load_workbook(original); ov = openpyxl.load_workbook(original, data_only=True)
    except Exception as e:
        r["judge_cannot_read_original"] = str(e)[:200]; return r
    try:
        c = openpyxl.load_workbook(copy); cv = openpyxl.load_workbook(copy, data_only=True)
        r["opens"] = True
    except Exception as e:
        r["opens"] = False; r["open_error"] = str(e)[:200]; return r
    skip_values = set(); edited = None
    if edit:
        info = json.load(open(edit))
        edited = (info["edit"]["sheet"], info["edit"]["ref"])
        skip_values = {tuple(k.rsplit("!", 1)) for k in info["expected"]} | {edited}
    # A workbook saved with no results (XlsxWriter, openpyxl, ExcelJS write 0 or nothing), or one that
    # asks Excel to recalculate on opening, has every
    # formula recalculated by the first edit, so its stored results are meant to change.
    formulas = [(ws.title, cell.coordinate) for ws in o.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f"]
    asks = b'fullCalcOnLoad="1"' in zipfile.ZipFile(original).read("xl/workbook.xml")
    placeholders = bool(edit) and bool(formulas) and (asks or all(ov[sheet][ref].value in (None, 0, "") for sheet, ref in formulas))
    cells = values_lost = formulas_lost = formats_lost = 0
    examples = []
    for ws in o.worksheets:
        if ws.title not in c.sheetnames: continue
        wc = c[ws.title]; wv = ov[ws.title]; wcv = cv[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and cell.style_id == 0: continue
                cells += 1
                ref = cell.coordinate
                other = wc[ref]
                if (ws.title, ref) != edited:
                    if formula_of(cell.value) != formula_of(other.value) and not (cell.data_type != "f" and same_value(cell.value, other.value)):
                        formulas_lost += 1
                        if len(examples) < 6: examples.append(("formula/value", ws.title, ref, str(cell.value)[:60], str(other.value)[:60]))
                    if (ws.title, ref) not in skip_values and not (placeholders and cell.data_type == "f") and not same_value(wv[ref].value, wcv[ref].value):
                        values_lost += 1
                        if len(examples) < 6: examples.append(("stored value", ws.title, ref, str(wv[ref].value)[:60], str(wcv[ref].value)[:60]))
                if look(cell) != look(other):
                    formats_lost += 1
                    if len(examples) < 6: examples.append(("format", ws.title, ref, str(look(cell))[:120], str(look(other))[:120]))
    r["cells"] = cells
    r["values_lost"] = values_lost; r["formulas_lost"] = formulas_lost; r["formats_lost"] = formats_lost
    sheet_problems = []
    if o.sheetnames != c.sheetnames: sheet_problems.append(f"sheets {o.sheetnames} -> {c.sheetnames}")
    for ws in o.worksheets:
        if ws.title not in c.sheetnames: continue
        wc = c[ws.title]
        if sorted(map(str, ws.merged_cells.ranges)) != sorted(map(str, wc.merged_cells.ranges)): sheet_problems.append(f"{ws.title}: merged cells")
        cols = {k: (round(v.width or 0, 4), v.hidden) for k, v in ws.column_dimensions.items() if v.width or v.hidden}
        cols2 = {k: (round(v.width or 0, 4), v.hidden) for k, v in wc.column_dimensions.items() if v.width or v.hidden}
        if cols != cols2: sheet_problems.append(f"{ws.title}: column widths")
        rows = {k: (v.height, v.hidden) for k, v in ws.row_dimensions.items() if v.height or v.hidden}
        rows2 = {k: (v.height, v.hidden) for k, v in wc.row_dimensions.items() if v.height or v.hidden}
        if rows != rows2: sheet_problems.append(f"{ws.title}: row heights")
        if ws.freeze_panes != wc.freeze_panes: sheet_problems.append(f"{ws.title}: frozen panes")
    r["layout_problems"] = sheet_problems
    kept = {}
    for ws in o.worksheets:
        if ws.title not in c.sheetnames: continue
        wc = c[ws.title]
        kept.setdefault("conditional formats", [0, 0])
        kept.setdefault("data validation", [0, 0])
        a, b = cf_rules(ws), cf_rules(wc)
        kept["conditional formats"][0] += len(a); kept["conditional formats"][1] += len([x for x in a if x in b])
        a, b = dv_rules(ws), dv_rules(wc)
        kept["data validation"][0] += len(a); kept["data validation"][1] += len([x for x in a if x in b])
    a, b = names(o), names(c)
    kept["defined names"] = [len(a), len([k for k, v in a.items() if b.get(k) == v])]
    for feature, pattern in FEATURES.items():
        a, b = part_hashes(original, pattern), part_hashes(copy, pattern)
        kept[feature] = [len(a), len([k for k in a if k in b]), len([k for k, v in a.items() if b.get(k) == v])]
    for feature, regex in RAW.items():
        kept[feature] = [raw_count(original, regex), raw_count(copy, regex)]
    r["kept"] = kept
    r["examples"] = examples
    return r

def judge_in_child(job):
    """Each file in a process of its own, at most two minutes: a workbook can hang or crash openpyxl."""
    original, copy, edit = job
    args = [sys.executable, os.path.abspath(__file__), "--one", original, copy] + ([edit] if edit else [])
    try:
        done = subprocess.run(args, capture_output=True, text=True, timeout=120)
        return json.loads(done.stdout.strip().splitlines()[-1])
    except Exception as e:
        return {"file": os.path.basename(copy), "judge_cannot_read_original": f"judge failed: {type(e).__name__}"}

if __name__ == "__main__":
    if sys.argv[1] == "--one":
        print(json.dumps(judge((sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else None)), default=str))
        sys.exit(0)
    originals, copies, out = sys.argv[1:4]
    jobs = []
    for name in sorted(os.listdir(copies)):
        if not name.endswith(".xlsx") or not os.path.exists(os.path.join(originals, name)): continue
        edit = os.path.join(copies, name[:-5] + ".json")
        jobs.append((os.path.join(originals, name), os.path.join(copies, name), edit if os.path.exists(edit) else None))
    with ThreadPoolExecutor(6) as pool:
        results = list(pool.map(judge_in_child, jobs))
    json.dump(results, open(out, "w"), indent=1, default=str)
    print(len(results), "judged")
