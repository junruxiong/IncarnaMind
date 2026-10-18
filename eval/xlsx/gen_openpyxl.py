"""Test workbooks written by openpyxl (MIT): its own dialect, formulas without stored results. Test data only."""
import sys, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
OUT = sys.argv[1]

def budget():
    wb = Workbook()
    ws = wb.active; ws.title = "Budget"
    ws["A1"] = "Department budget 2027"; ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1"); ws["A1"].alignment = Alignment(horizontal="center")
    heads = ["Line", "Q1", "Q2", "Q3", "Q4", "Year", "Share"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=3, column=c, value=h); cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.border = Border(bottom=Side(style="medium"))
    lines = ["Salaries", "Contractors", "Software", "Hardware", "Travel", "Training", "Office", "Marketing"]
    for r, name in enumerate(lines, 4):
        ws.cell(row=r, column=1, value=name)
        for c in range(2, 6):
            ws.cell(row=r, column=c, value=1000 * (r + c) + 250 * c).number_format = "#,##0"
        ws.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})").number_format = "#,##0"
        ws.cell(row=r, column=7, value=f"=F{r}/$F$12").number_format = "0.0%"
    ws.cell(row=12, column=1, value="Total").font = Font(bold=True)
    for c in range(2, 7):
        col = "ABCDEFG"[c - 1]
        ws.cell(row=12, column=c, value=f"=SUM({col}4:{col}11)").number_format = "#,##0"
    ws["A14"] = "Inflation"; ws["B14"] = 0.03; ws["B14"].number_format = "0.0%"
    ws["A15"] = "Next year"; ws["B15"] = "=F12*(1+Inflation)"; ws["B15"].number_format = "#,##0"
    ws["A16"] = "Over 30k"; ws["B16"] = '=COUNTIF(F4:F11,">30000")'
    ws["A17"] = "Biggest"; ws["B17"] = "=INDEX(A4:A11,MATCH(MAX(F4:F11),F4:F11,0))"
    ws["A3"].comment = Comment("Lines follow the chart of accounts", "Test")
    wb.defined_names["Inflation"] = DefinedName("Inflation", attr_text="Budget!$B$14")
    ws.conditional_formatting.add("F4:F11", CellIsRule(operator="greaterThan", formula=["30000"], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add("G4:G11", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"))
    ws.conditional_formatting.add("A4:A11", FormulaRule(formula=["$F4>35000"], font=Font(bold=True)))
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100000", showErrorMessage=True)
    ws.add_data_validation(dv); dv.add("B4:E11")
    ws.column_dimensions["A"].width = 16
    for col in "BCDEF": ws.column_dimensions[col].width = 11
    ws.freeze_panes = "B4"
    ch = BarChart(); ch.title = "Year by line"
    ch.add_data(Reference(ws, min_col=6, min_row=3, max_row=11), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=11))
    ws.add_chart(ch, "I3")
    s2 = wb.create_sheet("预算说明")
    s2["A1"] = "说明"; s2["B1"] = "本表为测试数据"
    s2["A2"] = "总额"; s2["B2"] = "=Budget!F12"; s2["B2"].number_format = '[$¥-804]#,##0'
    s2["A3"] = "日期"; s2["B3"] = datetime.date(2027, 3, 31); s2["B3"].number_format = 'yyyy"年"m"月"d"日"'
    s2["A4"] = "年份"; s2["B4"] = "=YEAR(B3)"
    s2["A5"] = "月末"; s2["B5"] = "=EOMONTH(B3,1)"; s2["B5"].number_format = "yyyy-mm-dd"
    s2["A6"] = "文字"; s2["B6"] = '=B1&"（"&TEXT(B2,"#,##0")&"）"'
    wb.save(os.path.join(OUT, "op-budget-预算.xlsx"))

def hr():
    wb = Workbook(); ws = wb.active; ws.title = "Staff"
    ws.append(["Name", "Team", "Start", "Salary", "Bonus %", "Bonus", "Years"])
    teams = ["Ops", "Sales", "Eng", "Finance"]
    for i in range(1, 41):
        r = i + 1
        ws.append([f"Person {i}", teams[i % 4], datetime.date(2015 + i % 10, 1 + i % 12, 1 + i % 27), 42000 + i * 900, (i % 5) / 100,
                   f"=ROUND(D{r}*E{r},0)", f'=DATEDIF(C{r},DATE(2027,1,1),"Y")'])
        ws.cell(row=r, column=3).number_format = "dd/mm/yyyy"; ws.cell(row=r, column=4).number_format = '"$"#,##0'
        ws.cell(row=r, column=5).number_format = "0%"
    ws["I1"] = "Team"; ws["J1"] = "Payroll"; ws["K1"] = "Average"; ws["L1"] = "Senior"
    for k, t in enumerate(teams, 2):
        ws[f"I{k}"] = t; ws[f"J{k}"] = f"=SUMIFS(D:D,B:B,I{k})"; ws[f"K{k}"] = f"=AVERAGEIFS(D$2:D$41,B$2:B$41,I{k})"
        ws[f"L{k}"] = f'=COUNTIFS(B$2:B$41,I{k},G$2:G$41,">=5")'
    ws["I7"] = "All"; ws["J7"] = "=SUM(J2:J5)"; ws["K7"] = "=AVERAGE(D2:D41)"; ws["L7"] = "=SUMPRODUCT((G2:G41>=5)*1)"
    ws["I9"] = "Median"; ws["J9"] = "=MEDIAN(D2:D41)"; ws["I10"] = "Stdev"; ws["J10"] = "=STDEV(D2:D41)"
    ws["I11"] = "Top"; ws["J11"] = "=LARGE(D2:D41,3)"; ws["I12"] = "Rank P7"; ws["J12"] = "=RANK(D8,D2:D41)"
    ws["I13"] = "Weekday"; ws["J13"] = "=WEEKDAY(C2)"; ws["I14"] = "Workdays"; ws["J14"] = "=NETWORKDAYS(C2,C3)"
    ws.auto_filter.ref = "A1:G41"
    wb.save(os.path.join(OUT, "op-staff.xlsx"))

budget(); hr(); print("ok")
