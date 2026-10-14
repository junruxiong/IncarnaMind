#!/usr/bin/env python3
"""
Writes the checks evaluation's own Documents (eval/checks/README.md, issue
#86): the ten documents someone else might send, to verify claim by claim
(received/), and the two web pages the Folder saves by their address (web/).
The rest of the Folder is the every-format set's Documents and one Chinese
Wikipedia article, already in the repository. From the repository root:

    python3 -m venv <venv>
    <venv>/bin/pip install reportlab==5.0.1 python-docx==1.2.0 python-pptx==1.0.2 openpyxl==3.1.5
    <venv>/bin/python -I eval/checks/make-fixtures.py

All text is synthetic, written for this project (see ATTRIBUTION.md). The
deterministic Office and PDF helpers are the every-format script's
(eval/retrieval/formats/make-fixtures.py), so the output is the same bytes
for the same library versions.
"""
import importlib.util
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate

HERE = Path(__file__).resolve().parent
RECEIVED = HERE / "received"
WEB = HERE / "web"


def every_format_script():
    sys.dont_write_bytecode = True
    path = HERE.parent / "retrieval" / "formats" / "make-fixtures.py"
    spec = importlib.util.spec_from_file_location("every_format_fixtures", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


FORMATS = every_format_script()


# ---------------------------------------------------------------------------
# Writers: a document is a title, a byline and blocks ("h", text) or ("p", text)
# ---------------------------------------------------------------------------


def write_pdf(path, language, title, byline, blocks):
    styles = FORMATS.pdf_styles(language)
    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        invariant=1,
        title="",
        author="",
        creator="",
        subject="",
    )
    flow = [Paragraph(title, styles["title"]), Paragraph(byline, styles["byline"])]
    for kind, text in blocks:
        flow.append(Paragraph(text, styles["h" if kind == "h" else "p"]))
    document.build(flow)


def write_docx(path, title, byline, blocks):
    document = Document()
    document.add_heading(title, level=0)
    document.add_paragraph(byline)
    for kind, text in blocks:
        if kind == "h":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    FORMATS.fix_core(document.core_properties)
    FORMATS.save_office(path, document.save)


def write_pptx(path, title, subtitle, slides):
    """slides: (title, bullets, notes or None)."""
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    cover.placeholders[1].text = subtitle
    for heading, bullets, notes in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = heading
        frame = slide.placeholders[1].text_frame
        for index, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    FORMATS.fix_core(deck.core_properties)
    FORMATS.save_office(path, deck.save)


def write_web_page(path, language, title, site, published, lede, sections):
    """A news page as a site serves it: navigation and a footer around the article."""
    body = "\n".join(
        f"      <h2>{FORMATS.esc(heading)}</h2>\n"
        + "\n".join(f"      <p>{FORMATS.esc(text)}</p>" for text in paragraphs)
        for heading, paragraphs in sections
    )
    path.write_text(
        f"""<!doctype html>
<html lang="{language}">
  <head>
    <meta charset="utf-8">
    <title>{FORMATS.esc(title)}</title>
    <meta property="og:site_name" content="{FORMATS.esc(site)}">
    <meta property="article:published_time" content="{published}">
  </head>
  <body>
    <nav><a href="/">{FORMATS.esc(site)}</a> · <a href="/news">News</a></nav>
    <article>
      <h1>{FORMATS.esc(title)}</h1>
      <p>{FORMATS.esc(lede)}</p>
{body}
    </article>
    <footer>{FORMATS.esc(site)} · synthetic page for IncarnaMind's evaluation</footer>
  </body>
</html>
""",
        encoding="utf-8",
    )


def h(text):
    return ("h", text)


def p(*sentences):
    return ("p", " ".join(sentences))


def p_zh(*sentences):
    return ("p", "".join(sentences))


# ---------------------------------------------------------------------------
# The documents to verify: five English, five Chinese
# ---------------------------------------------------------------------------

CONTRACTOR_REPORT = [
    h("1 Summary"),
    p(
        "This report covers the roof and insulation package from 1 to 31 March 2027.",
        "Roof works were programmed to start on 4 January 2027.",
        "Under our contract, the roof and insulation package is due to be completed in April 2027.",
        "Our contract for the roof and insulation package is valued at £1,240,000.",
    ),
    h("2 Progress"),
    p(
        "Forty-two operatives worked on site during March.",
        "Insulation to the reading room ceiling was 60 per cent complete by the end of March.",
        "Two working days were lost to high winds in the first week of March.",
        "The scaffolding was supplied and erected by Apex Access Ltd.",
    ),
    p(
        "The roof must be watertight before the lift installation can begin.",
        "Lifts and ramps are being installed by Hartley Building Ltd under the same contract.",
    ),
    h("3 Site and logistics"),
    p(
        "The site compound takes up half of the staff car park.",
        "Progress meetings are held every Tuesday at the council offices.",
        "All waste from the strip-out has been recycled.",
    ),
    h("4 Risks and costs"),
    p(
        "Asbestos removal in the boiler room ceiling was priced at £64,000 in the March 2026 survey.",
        "The bat survey found no roost under the eaves.",
        "Our site manager walks the eaves with an ecologist every week.",
    ),
    p(
        "The council's contingency is held at 15 per cent of the works cost.",
        "The total renovation budget is £3.2 million.",
        "The new roof membrane carries a 25-year manufacturer's warranty.",
        "The replacement windows for the reading room will be triple-glazed.",
    ),
    h("5 Next month"),
    p("We look forward to the next progress meeting and thank the library staff for their patience."),
]

PRESS_RELEASE = [
    p("DRAFT FOR APPROVAL. Prepared by Lanyard Communications for Harbour Roasters."),
    p(
        "Harbour Roasters' coffee subscription reached 9,420 paid subscribers at the end of September.",
        "Every bag is roasted no more than 48 hours before it is dispatched.",
        "Two roasting shifts run every weekday at the company's Manchester roastery.",
        "A new roastery in Leeds begins shipping on 3 November, halving delivery times in the north.",
    ),
    p(
        "In September's customer survey the service scored a net promoter score of 61.",
        "Customers named freshness as the feature they liked best.",
        "About 9 per cent of subscribers cancelled in their first month.",
    ),
    p(
        "The Office plan delivers four bags a month for £58.",
        "Subscribers in Scotland are served from a depot in Edinburgh, which opened in July.",
        "Referral codes have brought in more than 2,000 subscribers since launch.",
    ),
    p(
        "Harbour Roasters was founded in 2019 by two former baristas.",
        "All of its coffee is certified organic and Fairtrade.",
        "The boxes are fully compostable.",
        "Subscribers can pause their plan at any time from the app.",
        "Explorer subscribers receive a tasting card with every bag.",
        "Deliveries are carbon-neutral through the courier's offset programme.",
        "The company plans to open a coffee bar at its roastery next spring.",
    ),
    p("“We are delighted that so many people have joined us,” said the founders."),
    h("Notes to editors"),
    p("Images are available on request from the press office."),
]

STAFFING_PROPOSAL = [
    h("1 Your plan as we understand it"),
    p(
        "Your 2027 staffing plan sets a total budget of £50,929,200.",
        "Westbury, in the East region, plans for 19 nurses and 7 doctors.",
        "The plan budgets £438,000 for Ashford Vale.",
        "Marsh End, a South region clinic, is planned at 350 weekly hours.",
    ),
    p(
        "Hawthorn, in the West region, is planned with 12 nurses.",
        "Kingsmead's budget is £656,800.",
        "The staffing plan was approved by Dr Helen Okafor.",
        "The plan is due for review on 30 November 2026.",
    ),
    h("2 Equipment you have ordered"),
    p(
        "Sonaris Imaging is supplying six ultrasound scanners.",
        "The defibrillators are due for delivery in February 2027.",
        "The ECG machines come with a five-year service contract.",
        "The vaccine fridges will be installed by ColdChain Direct's own engineers.",
    ),
    h("3 Where we can help"),
    p(
        "Agency nurses currently fill 14 per cent of shifts across your network.",
        "The North region has the highest vacancy rate in your network.",
        "Each clinic must keep at least one doctor on site during opening hours.",
        "Your network employs 85 healthcare assistants.",
        "Cedar Park will open on Saturdays from January.",
    ),
    p(
        "We would be pleased to discuss this proposal at your convenience.",
        "Our rates are set out in the attached schedule.",
    ),
]

OPERATOR_SLIDES = [
    (
        "Station highlights",
        [
            "Gas Works, in Central, recorded 4,775 trips in September.",
            "Science Park recorded 5,750 trips in September.",
            "Abbey Road has 12 docks and averaged 9-minute trips.",
            "Albert Dock is in the Old Town district.",
        ],
        None,
    ),
    (
        "Station details",
        [
            "Bus Station, in the University Quarter, averaged 27-minute trips.",
            "Wool Hall has 32 docks.",
            "Trips from Canal Basin averaged 25.5 minutes.",
            "Ferry Terminal, in Old Town, has 22 docks.",
            "Town Hall station sits in the Central district.",
        ],
        None,
    ),
    (
        "Fleet and members",
        [
            "E-bikes made up 35 per cent of the fleet in September.",
            "Bikes were available at 92 per cent of stations at peak times.",
            "We added 1,200 annual members in September.",
        ],
        "Thank the transport team for their support this month.",
    ),
    (
        "Looking ahead",
        [
            "Twelve new stations will open in the Docklands next spring.",
            "Vandalism cost the scheme £14,000 in September.",
            "The Hospital station is used mostly by night-shift staff.",
            "Weekend trips were 40 per cent longer than weekday trips.",
        ],
        None,
    ),
]

ADVISORY_MEMO = [
    h("Findings"),
    p(
        "The study measured temperatures in six districts between June and August 2026.",
        "The highest surface temperature, 47.2 °C, was recorded in Old Market on 19 July.",
        "Eastgate's tree cover is 6 per cent, well below the city average of 19 per cent.",
        "Hillside was on average 4.3 °C cooler than Eastgate at noon.",
    ),
    p(
        "Heat-related hospital admissions in Eastgate doubled between 2020 and 2025.",
        "Older residents in Eastgate spend on average six hours a day indoors in summer.",
    ),
    h("Cooling measures"),
    p(
        "Cool roofs were the most expensive measure per square metre.",
        "Street trees cost £130 each to plant.",
        "Each misting station uses about 40 litres of water an hour.",
        "The city installed 20 misting stations in June 2026.",
    ),
    p(
        "Permeable paving also reduced surface flooding on the test streets.",
        "Cool roofs were first trialled in the city in 2019.",
        "Misting stations halved heat complaints at the bus stops where they were installed.",
    ),
    h("Next steps"),
    p(
        "A follow-up study with 80 indoor loggers will start in June 2027.",
        "The study was funded by a national research council grant.",
        "The city plans to plant 5,000 street trees by 2030.",
    ),
    p("We recommend the committee discuss these findings at its next meeting."),
]

CATERER_MEMO = [
    p_zh("致城东街道办事处："),
    h("一、基本情况"),
    p_zh(
        "桂花苑食堂由我公司运营，日均供餐260份。",
        "新村食堂日均供餐210份，满意度达到92%。",
        "老年居民每餐自付8元，其余部分由区财政补贴。",
        "试点期间各食堂平均每餐成本为16.8元。",
    ),
    p_zh(
        "食材统一由区农产品配送中心采购配送。",
        "配送中心的合同将于2027年6月到期。",
    ),
    h("二、运营情况"),
    p_zh(
        "增设第二个打饭窗口后，桂花苑食堂高峰排队时间缩短到8分钟左右。",
        "桂花苑片区的晚餐供应将于2027年3月开始。",
        "评估期内新村食堂实现盈利约42万元。",
        "新村食堂的志愿者送餐员有9人。",
    ),
    p_zh(
        "我公司厨房员工均持有有效健康证。",
        "桂花苑食堂每周推出两款新菜品。",
    ),
    h("三、下一步打算"),
    p_zh(
        "我公司已为新村食堂采购了一台新的蒸饭柜。",
        "新村食堂的租金每年为36万元。",
        "桂花苑食堂的厨房于2025年5月完成排烟改造。",
        "我公司计划在2027年投资50万元建设中央厨房。",
    ),
    p_zh("特此说明，请予审阅。"),
]

BUS_PRESS_RELEASE = [
    p_zh("江城融媒体中心讯（2026年10月12日）"),
    p_zh(
        "第三季度，江城公交集团纯电动公交车占比达到87%。",
        "第三季度准点率提升至95.3%。",
        "本季度新增低地板无障碍车辆15辆。",
        "1路日均客流18,600人次，准点率95.2%。",
    ),
    p_zh(
        "第二座充电站将于2027年1月投入使用，位于高铁西站。",
        "本季度共更换电池组16套。",
        "夜班线路N3将于11月1日开通。",
        "乘客支付方式中，公交卡占比最高。",
        "驾驶员疲劳监测系统已覆盖全部车辆。",
    ),
    p_zh(
        "目前全市公交线路总数达到186条。",
        "集团今年共招聘新驾驶员240名。",
        "9路将于明年改为纯电动双层巴士。",
        "百公里电耗比去年同期下降了4%。",
    ),
    p_zh(
        "老年人持卡乘车继续免费。",
        "高新区自动驾驶接驳线路将投入10辆无人驾驶小巴。",
        "第三季度共处理乘客投诉1,250件。",
    ),
    p_zh("集团表示，将继续提升服务质量。"),
]

SUPPLIER_STATEMENT = [
    p_zh("致：门店运营部"),
    h("一、销售情况"),
    p_zh(
        "第三季度门店合计销售额为50,525,225元。",
        "人民路广场店（嘉兴）第三季度同比增长22.0%。",
        "复兴路旗舰店第三季度销售额为1,050,573元。",
        "本季度销售数据截止到2026年10月15日。",
    ),
    h("二、库存情况"),
    p_zh(
        "萧山冷链仓有机鲜牛奶库存量为4200，高于3000的安全库存。",
        "鲜榨橙汁库存为2,320，高于安全库存。",
        "鲜榨橙汁的补货周期为2天。",
        "冷冻水饺存放在滨江中心仓。",
        "坚果礼盒的补货周期为14天。",
    ),
    h("三、供货说明"),
    p_zh(
        "我公司供应的有机鲜牛奶保质期为7天。",
        "第三季度我公司向各门店共配送鲜奶18万盒。",
        "双方约定的结算周期为月结30天。",
        "全麦吐司由我公司旗下烘焙工厂每日供货。",
    ),
    p_zh(
        "第四季度起，鲜奶供货价格上调3%。",
        "滨江中心仓于2026年8月完成扩建。",
        "手冲咖啡豆的主要产地为云南保山。",
    ),
    p_zh("如有疑问，请与我公司财务部联系。"),
]

NEWSPAPER_ARTICLE = [
    p_zh("本报讯"),
    p_zh(
        "2025年，江城市日最高气温超过35℃的天数达到41天。",
        "老城区正午地表温度最高达到48.6℃。",
        "老城区绿地率仅为12%。",
        "西湖区绿地率最高，达到38%。",
    ),
    p_zh(
        "北山区夏季平均气温比老城区低2℃。",
        "高温天老城区急诊就诊人数增加了三成。",
        "临江区居民中有一半以上家中没有空调。",
    ),
    h("降温措施"),
    p_zh(
        "喷雾降温设施每套每小时耗水约45升。",
        "行道树每棵种植成本为1,600元。",
        "屋顶绿化在老城区坡屋顶住宅上效果良好，已全面推广。",
    ),
    p_zh(
        "透水铺装还能减少暴雨时的路面积水。",
        "市政府已决定明年新建20处喷雾降温设施。",
    ),
    h("研究团队"),
    p_zh(
        "这项调查由江城市气象台完成。",
        "调查共花费经费120万元。",
        "研究团队共有12名成员。",
        "研究建议优先在老城区和临江区补种行道树。",
    ),
    p_zh("专家提醒市民高温天减少午间外出。"),
]

INSPECTION_REPORT = [
    h("一、制度核对"),
    p_zh(
        "仓库规定冷冻区温度不得高于零下18度。",
        "停电超过10分钟时须启动备用发电机。",
        "灭火器每月15日由安保部检查。",
        "供应商车辆须提前一天预约卸货时段。",
    ),
    p_zh(
        "夜班负责人为孙雅琴，交接时间为00:00。",
        "冷藏车车厢温度高于零下15度的不予接收。",
        "制冷机组每四小时巡检一次。",
        "门禁系统在门打开超过60秒后报警。",
        "按规定，外来司机可以进入冷库作业区协助卸货。",
    ),
    h("二、现场情况"),
    p_zh(
        "仓库冷冻区面积为3,200平方米。",
        "现场抽查的12台叉车中有2台电池老化。",
        "仓库已通过ISO 22000食品安全管理体系认证。",
        "备用发电机每月进行一次带载试运行。",
    ),
    p_zh(
        "温度记录保存期限为两年。",
        "仓库共安装了48个监控摄像头。",
        "叉车司机每年须参加一次安全培训。",
    ),
    h("三、结论"),
    p_zh("总体来看，仓库管理较为规范。"),
]


# ---------------------------------------------------------------------------
# The Folder's two web pages, saved by their address
# ---------------------------------------------------------------------------

COUNCIL_PAGE = {
    "title": "Roof contract signed for Riverside Library",
    "site": "Riverside Borough Council",
    "published": "2026-12-02",
    "lede": (
        "Riverside Borough Council has signed the contract for the roof and insulation works at Riverside "
        "Library with Hartley Building Ltd, the first of the renovation's four works packages."
    ),
    "sections": [
        (
            "What happens first",
            [
                "Work on the roof begins on 4 January 2027. Scaffolding will go up on the Wharf Street side of "
                "the building in the last week of December.",
                "Hartley Building Ltd will take on four local apprentices for the length of the contract.",
            ],
        ),
        (
            "The stained glass window",
            [
                "The reading room's stained glass window, made for the library's opening in 1968, will be "
                "taken out and restored by a specialist studio at a cost of £28,000.",
            ],
        ),
        (
            "Paying for the renovation",
            [
                "The £2.6 million renovation is paid for by a £1.1 million grant from the national library "
                "improvement fund and by the council's own capital programme.",
            ],
        ),
        (
            "Using the library during the works",
            [
                "While the north wing is closed, the mobile library van will stop outside the leisure centre "
                "every Tuesday and Friday from 10am to 4pm.",
                "Monthly updates are posted on the library noticeboard and on this website.",
            ],
        ),
        (
            "Reopening",
            [
                "The council expects the renovated library to reopen in full in November 2027, with a family "
                "open weekend.",
            ],
        ),
    ],
}

BUS_PAGE = {
    "title": "夜班线路N3将于12月1日开通",
    "site": "江城公交集团",
    "published": "2026-11-15",
    "lede": "江城公交集团今天宣布，夜班线路N3将于2026年12月1日开通，每天晚上10点至次日凌晨5点运营。",
    "sections": [
        (
            "线路走向",
            [
                "N3路由北站枢纽发车，经人民路、解放路和中山路，开往高铁西站，全程18.6公里，共设14个站点。",
                "夜间发车间隔为30分钟，单程票价2元。",
            ],
        ),
        ("车辆与安全", ["N3路全部使用纯电动公交车，每班车配一名安全员随车。"]),
        ("北站枢纽充电站", ["北站枢纽充电站计划于2027年1月投入使用，建成后设有40个充电桩。"]),
        (
            "乘客意见征集",
            ["集团在9月通过线上平台征集乘客意见，共收到建议1,850条，其中超过一半希望延长末班车时间。"],
        ),
    ],
}


def main():
    RECEIVED.mkdir(parents=True, exist_ok=True)
    WEB.mkdir(parents=True, exist_ok=True)
    write_pdf(
        RECEIVED / "Hartley Building progress report 3.pdf",
        "en",
        "Riverside Library roof and insulation works: progress report 3",
        "Hartley Building Ltd · for Riverside Borough Council · 31 March 2027",
        CONTRACTOR_REPORT,
    )
    write_docx(
        RECEIVED / "Harbour Roasters press release draft.docx",
        "Harbour Roasters coffee subscription: press release",
        "For release on 6 October 2026",
        PRESS_RELEASE,
    )
    write_pdf(
        RECEIVED / "Meridian Health Staffing proposal.pdf",
        "en",
        "Proposal: nurse and doctor cover for the clinic network, 2027",
        "Meridian Health Staffing · for the clinic network's workforce team · 2 November 2026",
        STAFFING_PROPOSAL,
    )
    write_pptx(
        RECEIVED / "Cityride September 2026 operator report.pptx",
        "Cityride bike share: September 2026 report",
        "Prepared by Cityride Operations for the city's transport team",
        OPERATOR_SLIDES,
    )
    write_docx(
        RECEIVED / "Northgate Advisory memo on the heat study.docx",
        "Memo: what the summer 2026 heat study means for the climate committee",
        "From Northgate Advisory to the city's climate committee, 14 October 2026",
        ADVISORY_MEMO,
    )
    write_docx(
        RECEIVED / "禾丰餐饮运营情况说明.docx",
        "关于社区食堂运营情况的说明",
        "禾丰餐饮公司 · 2026年10月20日",
        CATERER_MEMO,
    )
    write_pdf(
        RECEIVED / "江城公交第三季度新闻稿.pdf",
        "zh",
        "新闻稿：江城公交第三季度运营情况",
        "江城融媒体中心 · 2026年10月12日",
        BUS_PRESS_RELEASE,
    )
    write_pdf(
        RECEIVED / "浙北乳品季度对账说明.pdf",
        "zh",
        "第三季度供货与对账说明",
        "浙北乳品供应链公司 · 2026年10月18日",
        SUPPLIER_STATEMENT,
    )
    write_docx(
        RECEIVED / "江城晚报绿地降温报道.docx",
        "研究称：绿地越少的城区越热",
        "江城晚报 · 2026年10月16日",
        NEWSPAPER_ARTICLE,
    )
    write_pdf(
        RECEIVED / "冷链仓库第三方巡检报告.pdf",
        "zh",
        "冷链仓库现场巡检报告",
        "华安仓储检测服务公司 · 巡检日期2026年10月9日",
        INSPECTION_REPORT,
    )
    write_web_page(WEB / "riverside-roof-contract.html", "en", **COUNCIL_PAGE)
    write_web_page(WEB / "jiangcheng-n3-night-line.html", "zh-CN", **BUS_PAGE)
    for path in sorted([*RECEIVED.iterdir(), *WEB.iterdir()]):
        print(f"{path.stat().st_size:>8}  {path.parent.name}/{path.name}")


if __name__ == "__main__":
    main()
