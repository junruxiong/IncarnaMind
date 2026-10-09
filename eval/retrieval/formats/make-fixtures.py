#!/usr/bin/env python3
"""
Writes the mixed-format evaluation fixtures (eval/README.md, "Every format",
issue #70) into files/ next to this script: an English and a Chinese Word
file, deck, workbook, CSV, Markdown file, plain-text file and text PDF, each
with the hard places the Questions in ../formats.json ask about. From the
repository root:

    python3 -m venv <venv>
    <venv>/bin/pip install reportlab python-docx python-pptx openpyxl
    <venv>/bin/python -I eval/retrieval/formats/make-fixtures.py

All text is synthetic, written for this project (see ATTRIBUTION.md). The
scanned PDFs the set asks about are the Organize fixtures' image-only scans
(tests/fixtures/organize/files/), not written here.

Word footnotes are added as a footnotes part (python-docx has no call for
them); comments use python-docx's own. Chinese PDF text uses reportlab's
built-in STSong-Light CID font, not embedded (pdf.js maps it with its
CMaps). The output is the same bytes for the same library versions: fixed
document dates, fixed ZIP timestamps (also inside the workbooks charts
embed) and reportlab's invariant mode.
"""
import datetime
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    NextPageTemplate,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

OUT = Path(__file__).resolve().parent / "files"
FIXED = datetime.datetime(2026, 9, 1, 9, 0, 0)

pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


# ---------------------------------------------------------------------------
# Deterministic Office files
# ---------------------------------------------------------------------------


def normalized_zip(data):
    """An Office package's bytes with fixed timestamps, inside embedded packages too."""
    stamp = FIXED.strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = [(info.filename, archive.read(info.filename)) for info in archive.infolist()]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, part in items:
            if name.endswith("core.xml"):
                part = re.sub(
                    r"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)",
                    lambda m: m.group(1) + stamp + m.group(3),
                    part.decode("utf-8"),
                ).encode("utf-8")
            elif name.endswith(".xlsx"):
                part = normalized_zip(part)
            info = zipfile.ZipInfo(name, date_time=FIXED.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, part)
    return out.getvalue()


def save_office(path, save):
    buffer = io.BytesIO()
    save(buffer)
    path.write_bytes(normalized_zip(buffer.getvalue()))


def fix_core(properties):
    properties.created = FIXED
    properties.modified = FIXED
    properties.last_printed = FIXED
    properties.author = ""
    properties.last_modified_by = ""
    properties.revision = 1


# ---------------------------------------------------------------------------
# Word: sections, a table, footnotes and comments
# ---------------------------------------------------------------------------

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def esc(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def footnote_reference(paragraph, number):
    """A superscript footnote reference at the end of the paragraph."""
    paragraph._p.append(
        parse_xml(
            f'<w:r {W}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
            f'<w:footnoteReference w:id="{number}"/></w:r>'
        )
    )


def add_footnotes(document, notes):
    """The footnotes part, as Word writes it: two separators, then each note."""
    body = (
        f'<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        f'<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    )
    for number, text in enumerate(notes, 1):
        body += (
            f'<w:footnote w:id="{number}"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
            f'<w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> {esc(text)}</w:t></w:r></w:p></w:footnote>'
        )
    xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {W}>{body}</w:footnotes>'
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        xml.encode("utf-8"),
        document.part.package,
    )
    document.part.relate_to(part, RT.FOOTNOTES)


def write_docx(path, title, blocks):
    """
    blocks: ("p", text, footnote or None, comment or None), ("h1"|"h2", text),
    ("table", rows). Footnotes are numbered in order; a comment is anchored on
    the whole paragraph.
    """
    document = Document()
    document.add_heading(title, level=0)
    notes = []
    for block in blocks:
        kind = block[0]
        if kind in ("h1", "h2"):
            document.add_heading(block[1], level=1 if kind == "h1" else 2)
        elif kind == "p":
            _, text, footnote, comment = block
            paragraph = document.add_paragraph()
            run = paragraph.add_run(text)
            if footnote:
                notes.append(footnote)
                footnote_reference(paragraph, len(notes))
            if comment:
                document.add_comment(run, text=comment, author="Reviewer", initials="RV")
        elif kind == "table":
            rows = block[1]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    if notes:
        add_footnotes(document, notes)
    fix_core(document.core_properties)
    for comment in document.comments:
        comment._comment_elm.set(qn("w:date"), FIXED.strftime("%Y-%m-%dT%H:%M:%SZ"))
    save_office(path, document.save)


def p(text, footnote=None, comment=None):
    return ("p", text, footnote, comment)


WORD_EN = [
    p("Prepared for Riverside Borough Council by the libraries team, October 2026."),
    ("h1", "1 Background"),
    p("The Riverside Library opened in 1968 and last had a major refurbishment in 1994."),
    p(
        "Visits fell from 212,000 in 2019 to 164,000 in 2025, while the number of registered members stayed close to 31,000.",
        footnote="Visit counts come from the door sensors installed in 2017, which miss visitors who arrive with school groups.",
    ),
    p("The council approved the renovation on 14 March 2026 after a public consultation that drew 1,870 responses."),
    p(
        "The building has three floors and a basement archive that holds the borough's local history collection "
        "of about 48,000 items, including parish maps, photographs and the minute books of the old harbour board."
    ),
    p(
        "Energy use in 2025 was 386 megawatt hours, about twice the figure for a modern library of the same size, "
        "mostly because of the single-glazed reading room windows and a gas boiler installed in 1981."
    ),
    p(
        "Opening hours were cut from 58 to 46 hours a week in 2023 to save staff costs. Saturday afternoons and "
        "Wednesday evenings were the sessions lost, and both were among the busiest before the cut."
    ),
    ("h1", "2 Costs"),
    p("Table 1 lists the main works packages, their contractors and when each is due to finish."),
    (
        "table",
        [
            ["Package", "Contractor", "Cost", "Completion"],
            ["Roof and insulation", "Hartley Building Ltd", "£1,240,000", "June 2027"],
            ["Lifts and ramps", "Ascend Access", "£385,000", "March 2027"],
            ["Children's library fit-out", "Little Oak Interiors", "£212,500", "September 2027"],
            ["Café and community room", "Brookside Joinery", "£178,000", "October 2027"],
        ],
    ),
    p(
        "A contingency of 10 per cent of the works cost is held by the council's capital programme board.",
        footnote="The works cost excludes VAT, which the council recovers in full.",
        comment="Finance has asked to lower the contingency to 8 per cent once the roof tender is signed.",
    ),
    p(
        "The total budget is £2.6 million, of which £1.1 million comes from a national library improvement grant "
        "and the rest from the council's own capital programme."
    ),
    ("h2", "2.1 North wing"),
    p("Temporary closure starts on 1 February 2027 and lasts eleven weeks."),
    p("Staff move to the mobile library van during the works."),
    p("Returned books are collected at the leisure centre reception."),
    ("h2", "2.2 South wing"),
    p("Temporary closure starts on 6 September 2027 and lasts seven weeks."),
    p("Staff move to the town hall annexe during the works."),
    p("Returned books are collected at the Mill Lane post office."),
    ("h1", "3 Risks"),
    p(
        "Asbestos was found in the boiler room ceiling during the March 2026 survey; removal is priced at £46,000.",
        footnote="Survey by Northfield Environmental, report NE-2291, dated 18 March 2026.",
        comment="Check whether the asbestos removal needs the building to be empty for a fortnight.",
    ),
    p(
        "A bat roost under the east eaves may delay the roof works until the end of the breeding season.",
        comment="The bat survey must be repeated in May because the first one was done in winter.",
    ),
    p("The children's library stays open throughout, using the side entrance on Wharf Street."),
    p(
        "Steel prices could add up to £90,000 to the roof package if the tender slips past January, because the "
        "contractor's quotation is only held for ninety days."
    ),
    ("h1", "4 Programme"),
    p(
        "Works start on 4 January 2027 with the roof, which must be watertight before the lifts are installed. "
        "The main contractor's site compound will take half of the staff car park until December 2027."
    ),
    p(
        "Progress meetings are held every Tuesday at the council offices, and a public update is posted on the "
        "library noticeboard and the council website on the first Monday of each month."
    ),
    p(
        "The local history archive is packed and moved to the county record office for the whole programme; "
        "researchers can book visits there on Thursdays."
    ),
    ("h1", "5 Consultation findings"),
    p(
        "The most requested change was longer opening hours, named by 62 per cent of respondents. Quiet study "
        "space came second, and free Wi-Fi that works on every floor came third."
    ),
    p(
        "Parents asked for a buggy park and a baby-changing room on the ground floor; both are now in the design. "
        "Older residents asked to keep the large-print section on the ground floor, near the entrance."
    ),
    p(
        "Fewer than 1 in 10 respondents asked for a café. It was kept because its rent will pay the running costs "
        "of the community room, which local groups can then book for free."
    ),
    ("h1", "6 Accessibility"),
    p(
        "The new lift serves all three floors and the basement, with a car large enough for a wheelchair user and "
        "a companion. Hearing loops will be fitted at the help desk, in the community room and in the children's "
        "library."
    ),
    p(
        "Two accessible toilets replace the single one on the first floor, and the front steps get a handrail on "
        "both sides and a ramp at a gradient of 1 in 15."
    ),
    ("h1", "7 Next steps"),
    p(
        "The cabinet will be asked to release the second half of the budget at its meeting on 12 November 2026. "
        "A temporary library manager will be appointed for the closure periods, paid from the contingency."
    ),
]

WORD_ZH = [
    p("本报告由城东街道办事处委托编写，评估期为2025年7月至2026年6月。"),
    ("h1", "一、项目背景"),
    p("试点共设立6家社区食堂，主要服务60岁以上的老年居民。"),
    p(
        "评估期内累计供餐41.6万份，日均供餐约1,140份。",
        footnote="供餐数量以收银系统记录为准，不含志愿者送餐上门的份数。",
    ),
    p("老年居民每餐自付8元，其余部分由区财政补贴。"),
    p("城东街道60岁以上老年居民约2.3万人，其中独居老人约3,100人，主要分布在老街和新村两个片区。"),
    p("试点前，老人主要依靠子女送餐或自己做饭。入户调查显示，有27%的独居老人每天只吃两顿饭，晚餐多以剩饭剩菜为主。"),
    ("h1", "二、运营数据"),
    p("表1列出了各食堂的运营单位、日均供餐量和居民满意度。"),
    (
        "table",
        [
            ["食堂", "运营单位", "日均供餐（份）", "满意度"],
            ["桂花苑食堂", "禾丰餐饮公司", "260", "94%"],
            ["滨江食堂", "邻里膳坊", "215", "89%"],
            ["老街食堂", "福寿团餐", "180", "91%"],
            ["新村食堂", "禾丰餐饮公司", "150", "86%"],
        ],
    ),
    p(
        "食材采购统一由区农产品配送中心负责。",
        footnote="配送中心每周一、三、五送货，冷链车辆由区商务局提供。",
        comment="配送中心的合同将于2026年12月到期，需提前续签。",
    ),
    p("各食堂平均每餐成本为18.6元，其中食材成本约占62%，人工成本约占27%。"),
    p("午餐高峰集中在11点至12点半。桂花苑食堂高峰时段排队时间曾超过20分钟，增设第二个打饭窗口后缩短到8分钟左右。"),
    ("h2", "（一）桂花苑片区"),
    p("下一阶段计划增加晚餐供应，预计于2026年11月开始，新增厨师2名。"),
    p("片区内另设一处助餐点，位于社区服务中心一楼。"),
    p("片区现有志愿者送餐员12人，主要为行动不便的老人送餐上门。"),
    ("h2", "（二）滨江片区"),
    p("下一阶段计划增加晚餐供应，预计于2027年3月开始，新增厨师1名。"),
    p("片区内另设一处助餐点，位于滨江公园东门。"),
    p("片区现有志愿者送餐员7人，周末由社区青年志愿者队补充。"),
    ("h1", "三、问题与建议"),
    p(
        "部分老人反映菜品偏咸，营养师已将每餐用盐量控制在4克以内。",
        footnote="用盐量标准参照《中国居民膳食指南（2022）》的建议。",
        comment="建议在下次评估中加入对低盐菜品满意度的单独统计。",
    ),
    p(
        "建议2027年将试点扩大到10家，并引入线上预订。",
        comment="线上预订需要先确认老年居民使用智能手机的比例。",
    ),
    p("部分食堂场地面积不足，老街食堂只有38个座位，高峰时段有老人只能打包回家。"),
    p("建议与社区卫生服务中心合作，为患有糖尿病、高血压等慢性病的老人提供定制餐。"),
    ("h1", "四、经费情况"),
    p("试点总投入1,260万元，其中区财政补贴占71%，街道配套资金占18%，其余来自企业和个人捐赠。"),
    p("每家食堂的一次性改造费用平均为96万元，主要用于厨房排烟、消防和无障碍设施改造。"),
    p("评估期内运营亏损最大的是新村食堂，全年亏损约42万元，主要原因是就餐人数少、租金高。"),
    ("h1", "五、居民反馈"),
    p("问卷共回收有效样本1,516份，其中80岁以上老人占31%。"),
    p("最受欢迎的菜品是红烧鱼块和清炒时蔬，最常被提到的意见是希望增加粗粮和软烂易嚼的菜。"),
    p("有居民希望周末也能正常供餐，目前各食堂周日停业。"),
    ("h1", "六、下一步工作"),
    p("2026年第四季度完成新村食堂的场地扩建，座位增加到80个。"),
    p("在各食堂引入刷脸支付，方便没有智能手机的老人就餐。"),
]


# ---------------------------------------------------------------------------
# PowerPoint: slide text, a table, charts from their caches, speaker notes
# ---------------------------------------------------------------------------


def write_pptx(path, slides):
    deck = Presentation()
    for spec in slides:
        kind = spec["kind"]
        if kind == "cover":
            slide = deck.slides.add_slide(deck.slide_layouts[0])
            slide.shapes.title.text = spec["title"]
            slide.placeholders[1].text = spec["subtitle"]
        elif kind == "bullets":
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = spec["title"]
            frame = slide.placeholders[1].text_frame
            for index, bullet in enumerate(spec["bullets"]):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = bullet
        elif kind == "table":
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = spec["title"]
            rows = spec["rows"]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(0.6), Inches(1.7), Inches(8.8), Inches(0.45) * len(rows)
            )
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    frame = shape.table.cell(r, c).text_frame
                    frame.text = value
                    frame.paragraphs[0].font.size = Pt(16)
        elif kind == "chart":
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = spec["title"]
            data = CategoryChartData()
            data.categories = spec["categories"]
            data.add_series(spec["series"], spec["values"])
            chart_type = XL_CHART_TYPE.PIE if spec.get("pie") else XL_CHART_TYPE.COLUMN_CLUSTERED
            chart = slide.shapes.add_chart(
                chart_type, Inches(0.8), Inches(1.6), Inches(8.4), Inches(5.2), data
            ).chart
            chart.has_title = True
            chart.chart_title.text_frame.text = spec["chart_title"]
            chart.has_legend = bool(spec.get("pie"))
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.RIGHT
                chart.legend.include_in_layout = False
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    fix_core(deck.core_properties)
    save_office(path, deck.save)


DECK_EN = [
    {
        "kind": "cover",
        "title": "Coffee Subscription Launch Review",
        "subtitle": "Harbour Roasters · Board update, October 2026",
        "notes": "Thank the operations team before starting.",
    },
    {
        "kind": "bullets",
        "title": "Launch results",
        "bullets": [
            "Paid subscribers reached 9,420 by the end of September.",
            "Average order value rose to £14.80 per delivery.",
            "Cancellations in the first month stayed below 6 per cent.",
        ],
        "notes": "If asked: the Leeds roastery starts shipping on 3 November, which halves delivery times in the north.",
    },
    {
        "kind": "bullets",
        "title": "What customers told us",
        "bullets": [
            "Net promoter score was 52 in the September survey.",
            "Freshness was the most praised feature, named in one comment in three.",
            "Most complaints were about boxes too big for a letterbox.",
            "Requests for a decaf blend came from 8 per cent of respondents.",
        ],
        "notes": "The survey had 1,214 responses, a response rate of 13 per cent.",
    },
    {
        "kind": "table",
        "title": "Plans by tier",
        "rows": [
            ["Tier", "Bags per month", "Price", "Subscribers"],
            ["Explorer", "1", "£12.50", "4,180"],
            ["Regular", "2", "£22.00", "3,960"],
            ["Office", "6", "£58.00", "1,280"],
        ],
        "notes": "The Office tier is sold mainly through two coworking chains.",
    },
    {
        "kind": "bullets",
        "title": "Operations",
        "bullets": [
            "Beans are roasted no more than 48 hours before dispatch.",
            "Courier costs average £2.35 per parcel, down from £2.90 at launch.",
            "Two roasting shifts run on weekdays at the Bristol roastery.",
            "Late deliveries fell to 2.1 per cent of parcels in September.",
        ],
        "notes": "The courier contract comes up for renewal in February; two other couriers have quoted.",
    },
    {
        "kind": "chart",
        "title": "Growth since launch",
        "chart_title": "New subscribers per month",
        "series": "New subscribers",
        "categories": ["Apr", "May", "Jun", "Jul", "Aug", "Sep"],
        "values": [820, 1140, 1560, 1710, 1930, 2260],
        "notes": "The jump in September follows the podcast sponsorship that started on 28 August.",
    },
    {
        "kind": "bullets",
        "title": "Marketing",
        "bullets": [
            "The podcast sponsorship cost £18,000 for eight weeks.",
            "Referral codes brought in 1,310 subscribers since launch.",
            "Paid social adverts were paused in August after costs per sign-up doubled.",
        ],
    },
    {
        "kind": "chart",
        "title": "Where subscribers live",
        "chart_title": "Subscribers by region (%)",
        "series": "Share of subscribers",
        "categories": ["North", "Midlands", "South", "Scotland"],
        "values": [34, 22, 31, 13],
        "pie": True,
        "notes": "Scotland is served from the Glasgow depot, which opened in July.",
    },
    {
        "kind": "bullets",
        "title": "Risks",
        "bullets": [
            "Green coffee prices have risen 22 per cent since March.",
            "One cooperative in Colombia supplies 40 per cent of our beans.",
            "Card payment failures cost about 3 per cent of renewals each month.",
        ],
        "notes": "We are trialling a second supplier, from Peru, from December.",
    },
    {
        "kind": "bullets",
        "title": "Next quarter",
        "bullets": [
            "Launch a decaf tier in November.",
            "Reach 12,000 paid subscribers by the end of December.",
            "Open gift subscriptions before Black Friday.",
        ],
        "notes": "The board asked for a cash-flow forecast before it approves the decaf tier.",
    },
]

DECK_ZH = [
    {
        "kind": "cover",
        "title": "新能源公交季度运营汇报",
        "subtitle": "江城公交集团 · 2026年第三季度",
        "notes": "开场先感谢各车队的配合。",
    },
    {
        "kind": "bullets",
        "title": "本季度要点",
        "bullets": [
            "纯电动公交车占比达到87%。",
            "百公里平均电耗为62.4千瓦时。",
            "准点率提升至93.5%。",
        ],
        "notes": "如被问到：第二座充电站将于2027年1月投入使用，位于北站枢纽。",
    },
    {
        "kind": "bullets",
        "title": "乘客满意度",
        "bullets": [
            "乘客满意度调查得分为86.7分。",
            "投诉最多的是早晚高峰车厢拥挤。",
            "本季度新增低地板无障碍车辆15辆。",
        ],
        "notes": "本次调查共回收问卷3,420份。",
    },
    {
        "kind": "table",
        "title": "各线路情况",
        "rows": [
            ["线路", "车辆数", "日均客流", "准点率"],
            ["1路", "42", "18,600", "95.2%"],
            ["9路", "28", "11,300", "91.8%"],
            ["26路", "16", "5,450", "93.0%"],
        ],
        "notes": "26路下月起延伸到高铁西站。",
    },
    {
        "kind": "bullets",
        "title": "车辆维护",
        "bullets": [
            "动力电池平均健康度为91%。",
            "本季度更换电池组6套，均在质保期内。",
            "故障车辆平均停运时间缩短至1.8天。",
        ],
        "notes": "电池质保期为8年或60万公里，以先到者为准。",
    },
    {
        "kind": "chart",
        "title": "充电量",
        "chart_title": "各月充电量（万千瓦时）",
        "series": "充电量",
        "categories": ["7月", "8月", "9月"],
        "values": [118, 126, 121],
        "notes": "8月用电高峰时段夜间充电比例达到78%。",
    },
    {
        "kind": "chart",
        "title": "乘客支付方式",
        "chart_title": "支付方式占比（%）",
        "series": "占比",
        "categories": ["手机扫码", "公交卡", "现金"],
        "values": [64, 29, 7],
        "pie": True,
        "notes": "现金支付比例比去年下降了5个百分点。",
    },
    {
        "kind": "bullets",
        "title": "安全运营",
        "bullets": [
            "本季度未发生一般及以上等级的责任事故。",
            "驾驶员疲劳监测系统已覆盖全部车辆。",
            "夜间末班车增加安全员随车，覆盖12条线路。",
        ],
    },
    {
        "kind": "bullets",
        "title": "下季度计划",
        "bullets": [
            "开通夜班线路N3。",
            "完成剩余13%燃油公交车的替换招标。",
            "在高新区试点自动驾驶接驳线路。",
        ],
        "notes": "N3路预计12月1日开通，运营时间为晚上10点至次日凌晨5点。",
    },
]


# ---------------------------------------------------------------------------
# Excel and CSV: labelled cells, long tables (a header repeated), a second sheet
# ---------------------------------------------------------------------------

CLINICS_EN = [
    "Ashford Vale", "Barnsfield", "Brook Green", "Castlegate", "Cedar Park", "Chalk Hill",
    "Cliffside", "Copper Lane", "Crossways", "Deanfield", "Eastmoor", "Elm Row",
    "Fairhaven", "Fernbank", "Fieldgate", "Foxley", "Glenmoor", "Greystone",
    "Harbourside", "Hawthorn", "Heathcote", "Highbury Rise", "Hollins", "Ivy Bridge",
    "Kingsmead", "Lark Rise", "Linden", "Longmead", "Lowfield", "Maple Cross",
    "Marsh End", "Meadowbank", "Millbrook", "Northgate", "Oakfield", "Old Quay",
    "Orchard Way", "Parkside", "Pine Hollow", "Queensway", "Redcliffe", "Riverside",
    "Rookwood", "Saltmarsh", "Sandown", "Southfields", "Stonebridge", "Thornbury",
    "Upper Weald", "Westbury", "Willowbank", "Woodlands",
]
REGIONS_EN = ["North", "East", "South", "West"]


def clinic_rows():
    rows = []
    for index, name in enumerate(CLINICS_EN):
        nurses = 6 + (index * 7) % 15
        doctors = 2 + (index * 5) % 6
        hours = 40 * (nurses + doctors) + (index * 13) % 40
        budget = 41_000 * nurses + 96_000 * doctors + (index * 1_700) % 9_000
        rows.append([name, REGIONS_EN[index % 4], nurses, doctors, hours, budget])
    return rows


STORES_ZH = [
    prefix + suffix
    for prefix in [
        "人民路", "解放路", "中山路", "建设路", "和平路", "胜利路", "文化路", "新华路",
        "滨河路", "青年路", "延安路", "长江路", "黄河路", "复兴路", "光明路",
    ]
    for suffix in ["店", "二店", "旗舰店", "社区店", "广场店"]
]
CITIES_ZH = ["杭州", "宁波", "温州", "绍兴", "嘉兴"]


def store_rows():
    rows = []
    for index, name in enumerate(STORES_ZH):
        sales = 380_000 + (index * 37_919) % 640_000
        growth = ((index * 29) % 41 - 12) / 100
        basket = 58 + (index * 11) % 70 + 0.5 * (index % 2)
        rows.append([name, CITIES_ZH[index % 5], sales, growth, basket])
    return rows


def write_xlsx(path, sheets):
    """sheets: (name, rows, formats by column); a cell value of datetime.date is a date."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows, formats in sheets:
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(list(row))
        for column, number_format in formats.items():
            for cell in sheet[column]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = number_format
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, datetime.datetime):
                    cell.number_format = "yyyy-mm-dd"
    workbook.properties.created = FIXED
    workbook.properties.modified = FIXED
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    save_office(path, workbook.save)


def staffing_workbook():
    rows = clinic_rows()
    total = sum(row[5] for row in rows)
    staffing = (
        [["Clinic Staffing Plan 2027"], ["Clinic", "Region", "Nurses", "Doctors", "Weekly hours", "Budget"]]
        + rows
        + [
            [],
            ["Total budget", total],
            ["Approved by", "Dr Helen Okoro"],
            ["Review date", datetime.datetime(2026, 11, 30)],
        ]
    )
    equipment = [
        ["Item", "Supplier", "Unit price", "Units ordered", "Delivery"],
        ["ECG machine", "Cardiotech Supplies", 4850, 12, "January 2027"],
        ["Defibrillator", "Lifeline Medical", 1390, 26, "December 2026"],
        ["Ultrasound scanner", "Sonaris Imaging", 18200, 4, "March 2027"],
        ["Examination couch", "Wardrobe Healthcare", 640, 40, "November 2026"],
        ["Blood pressure monitor", "Cardiotech Supplies", 85, 120, "November 2026"],
        ["Vaccine fridge", "ColdChain Direct", 2310, 9, "February 2027"],
    ]
    return [
        ("Staffing", staffing, {"F": '"£"#,##0', "B": '"£"#,##0'}),
        ("Equipment", equipment, {"C": '"£"#,##0'}),
    ]


def sales_workbook():
    rows = store_rows()
    total = sum(row[2] for row in rows)
    sales = (
        [["2026年第三季度门店销售"], ["门店", "城市", "销售额", "同比增长", "客单价"]]
        + rows
        + [
            [],
            ["合计销售额", total],
            ["制表人", "李晓雯"],
            ["数据截止日期", datetime.datetime(2026, 9, 30)],
        ]
    )
    stock = [
        ["商品", "仓库", "库存量", "安全库存", "补货周期（天）"],
        ["有机鲜牛奶", "萧山冷链仓", 4200, 3000, 2],
        ["全麦吐司", "滨江中心仓", 1850, 1200, 1],
        ["冷冻水饺", "萧山冷链仓", 6300, 2500, 7],
        ["手冲咖啡豆", "余杭常温仓", 940, 600, 14],
        ["鲜榨橙汁", "滨江中心仓", 1320, 1500, 2],
        ["坚果礼盒", "余杭常温仓", 2780, 800, 30],
    ]
    return [
        ("销售", sales, {"C": '"¥"#,##0', "B": '"¥"#,##0', "D": "0.0%", "E": "0.0"}),
        ("库存", stock, {}),
    ]


STATIONS_EN = [
    "Abbey Road", "Albert Dock", "Arch Street", "Bank Square", "Bell Lane", "Bishops Gate",
    "Boat House", "Bridge End", "Bus Station", "Canal Basin", "Castle Hill", "Cathedral",
    "Chapel Walk", "Church Green", "City Library", "Clock Tower", "College Road", "Corn Exchange",
    "Court House", "Cricket Ground", "Dock Gates", "Exchange Place", "Ferry Terminal", "Fire Station",
    "Fish Market", "Gas Works", "Grain Store", "Guild Hall", "Harbour Steps", "Hay Market",
    "High Street", "Hospital", "Ice Rink", "Iron Bridge", "King's Yard", "Lantern Square",
    "Leisure Centre", "Lock Keeper", "Market Cross", "Mill Race", "Museum", "New Quay",
    "Old Brewery", "Opera House", "Park Gates", "Pier Head", "Post Office", "Quayside",
    "Railway Arch", "Rope Walk", "Science Park", "Ship Yard", "Silk Mill", "Stadium",
    "Station Approach", "Swan Wharf", "Tannery", "Theatre", "Tower Bridge", "Town Hall",
    "Tram Depot", "University", "Velodrome", "Viaduct", "Water Tower", "Weavers Row",
    "West Pier", "Wool Hall", "Zoo Gate",
]
DISTRICTS_EN = ["Central", "Docklands", "Old Town", "University Quarter", "Riverside"]


def stations_csv():
    lines = ["station_id,station_name,district,docks,trips_september,avg_trip_minutes"]
    for index, name in enumerate(STATIONS_EN):
        docks = 12 + (index * 5) % 25
        trips = 900 + (index * 467) % 3900
        minutes = 9 + (index * 7) % 19 + 0.5 * (index % 2)
        lines.append(f"BS-{1001 + index},{name},{DISTRICTS_EN[index % 5]},{docks},{trips},{minutes:.1f}")
    return "\n".join(lines) + "\n"


ESTATES_ZH = [
    prefix + suffix
    for prefix in ["阳光", "翠湖", "金桂", "碧水", "锦绣", "枫林", "银杏", "明珠", "紫荆", "春晓", "海棠", "兰亭"]
    for suffix in ["花园", "苑", "家园", "新村", "雅居", "名邸"]
]
STREETS_ZH = ["东湖街道", "西溪街道", "南岸街道", "北山街道"]


def sorting_csv():
    lines = ["小区,街道,户数,9月分类准确率,厨余垃圾（吨）,督导员"]
    surnames = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张"
    given = ["建华", "秀英", "志强", "丽娟", "海燕", "国庆", "晓东", "春梅", "文斌", "雪芳", "立新", "美玲"]
    for index, name in enumerate(ESTATES_ZH):
        households = 320 + (index * 53) % 1400
        accuracy = 78 + (index * 7) % 21
        waste = 18 + (index * 13) % 70 + 0.5 * (index % 2)
        person = surnames[index % len(surnames)] + given[(index * 5) % len(given)]
        lines.append(f"{name},{STREETS_ZH[index % 4]},{households},{accuracy}%,{waste:.1f},{person}")
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Markdown and plain text: a table, lists, a code block
# ---------------------------------------------------------------------------

MARKDOWN_EN = """# Field Kit Setup Guide

This guide is for survey teams taking the standard field kit out for more than one night. Read it before your first trip and keep a printed copy in Bag A.

## Packing list

| Item | Weight | Packed in | Notes |
|---|---|---|---|
| Tent (two-person) | 2.1 kg | Bag A | Pegs in the side pocket |
| Water filter | 0.4 kg | Bag B | Replace the cartridge after 1,000 litres |
| Satellite messenger | 0.15 kg | Chest pouch | Charge to 100% the night before |
| Stove and fuel | 0.9 kg | Bag B | Fuel canisters can't fly in hold luggage |
| First aid kit | 0.6 kg | Bag A | Restock after every trip |

Bag A should weigh under 9 kg when packed; Bag B under 7 kg.

## Before you leave

1. Register the trip plan with the base office at least 48 hours ahead.
2. Test the satellite messenger by sending a check-in on channel 7.
3. Photograph the vehicle's fuel gauge and mileage.
4. Leave the spare vehicle key in the key safe at the depot.

Things to check on the weather forecast:

- Wind above 50 km/h on exposed ridges means the trip is postponed.
- Thunderstorms within 24 hours mean no camping above the tree line.
- A river level above 1.2 m at the Hollin gauge closes the ford.

## Syncing data

The kit's tablet syncs survey records to the base server whenever it has a signal. Its settings live in `sync.ini`:

```ini
[sync]
server = sync.fieldkit.example
port = 8443
interval_minutes = 20
retry_limit = 5
compress = true
```

If a sync fails five times in a row, the tablet stores records locally until the next manual sync.

Survey records are exported as CSV with one row per observation. The first line of every export is the header:

```csv
site_id,observed_at,observer,species_code,count,notes
```

Times are in UTC, written as ISO 8601. Species codes follow the regional four-letter list, so a curlew is CURL and a lapwing is LAPW.

## Charging and power

The kit runs on two power banks and a folding solar panel. In cloud the panel gives about a third of its rated output, so plan on the power banks alone for overcast trips.

| Device | Battery | Lasts | Charge from |
|---|---|---|---|
| Survey tablet | 7,600 mAh | 9 hours of fieldwork | Power bank 1 |
| Satellite messenger | 2,000 mAh | 4 days with 10-minute tracking | Power bank 2 |
| Head torch | 1,200 mAh | 6 hours on the medium setting | Either bank |

Keep the power banks inside your sleeping bag on cold nights: below freezing they can lose half their charge overnight.

## Radio procedure

Call the base office at 08:00 and 18:00 every day you are out, even when there is nothing to report. If a call is missed, base tries again after thirty minutes and then raises the alarm after a second missed call. Use plain language and give your grid reference to the nearest hundred metres.

## In an emergency

- Press SOS on the satellite messenger and keep it switched on.
- Stay with the vehicle or the tent unless it is unsafe to do so.
- Give first aid, then send a message with the number of people injured.

## After the trip

Dry the tent before it goes back into storage, and log any broken equipment in the kit register within two days. Return the satellite messenger to the charging shelf in the store room, not to a desk drawer, so that the next team finds it charged.
"""

MARKDOWN_ZH = """# 实验室数据管理规范

本规范适用于生命科学楼三层全部课题组，自2026年9月1日起执行。

## 存储位置与保留期限

| 数据类型 | 存储位置 | 保留期限 | 负责人 |
|---|---|---|---|
| 原始测序数据 | 冷存储阵列B | 10年 | 陈思远 |
| 显微图像 | 影像服务器 | 5年 | 王丽华 |
| 实验记录本扫描件 | 档案系统 | 永久 | 赵敏 |
| 分析中间文件 | 计算集群临时盘 | 90天 | 各课题组 |

个人电脑不得作为任何数据的唯一存放位置。

## 数据提交流程

1. 实验结束后24小时内上传原始数据。
2. 上传前用校验工具生成SHA-256摘要，并与数据一同提交。
3. 由课题组长在系统中确认后归档。
4. 归档后的数据只读，修改须另建新版本。

提交时常见的问题：

- 文件名含有空格或中文括号，会导致校验失败。
- 单个文件超过200GB时，须先联系平台工程师分卷。
- 涉及人类样本的数据须附伦理审批编号。

## 备份脚本

平台每晚自动备份，配置如下：

```bash
# 每晚自动备份
BACKUP_TARGET=/mnt/archive/lab
KEEP_DAYS=45
BANDWIDTH_LIMIT=80M
rsync -a --delete --bwlimit="$BANDWIDTH_LIMIT" /data/lab/ "$BACKUP_TARGET"
```

备份失败时，系统会在次日早上8点前向值班工程师发送提醒。

## 文件命名规则

所有上传的数据文件须按以下格式命名，各部分之间用下划线连接：

```text
课题编号_样本编号_实验日期_版本号
例：LS2026-07_S0142_20260915_v2
```

实验日期一律写成八位数字，版本号从v1开始递增。

## 访问权限

| 角色 | 原始数据 | 分析结果 | 审批人 |
|---|---|---|---|
| 课题组成员 | 只读 | 读写 | 课题组长 |
| 合作单位人员 | 无 | 只读 | 平台主任 |
| 平台工程师 | 读写 | 读写 | 平台主任 |

权限每半年复核一次，离组人员的权限在离组当天关闭。

## 离职交接

人员离职前须完成数据交接，并由课题组长签字确认。交接清单须列出全部数据的存放路径和对应的实验记录本页码。
"""

TEXT_EN = """NIGHT SHIFT HANDBOOK - MEDICAL WARDS
Revised September 2026. Keep this copy at the nurses' station.

1. WHO IS ON

Ward      Lead nurse          Handover   Beds
Ward 3    Priya Nandakumar    19:30      24
Ward 5    Tomasz Wieczorek    19:45      18
Ward 7    Aileen Doherty      20:00      30
Ward 9    Samuel Osei         20:15      22

The site manager covers all four wards from 22:00 and carries the red bleep.

2. HANDOVER CHECKLIST

At every handover, the incoming nurse and the outgoing nurse together:
  - count the controlled drugs and sign the register,
  - check the resuscitation trolley's seal number against the log,
  - list every patient on 15-minute observations,
  - confirm which patients are nil by mouth after midnight.

Do not start the drug round until the checklist is signed.

3. ESCALATION

Call the on-call registrar for a NEWS2 score of 5 or more, or a rise of 3 in one hour.
Call the outreach team for any patient who needs oxygen above 6 litres a minute.
For a cardiac arrest, dial 2222 and say the ward and bed number twice.

4. PAGER SETUP

New pagers are set up by the switchboard. If one needs resetting, the settings are:

    PAGER_GROUP=nights-medical
    ESCALATE_AFTER=10
    FALLBACK=switchboard
    QUIET_HOURS=none

A page that is not acknowledged within ESCALATE_AFTER minutes goes to the fallback.

5. BREAKS

Each nurse takes a 45-minute break between 01:00 and 04:00; breaks are agreed at the 22:00 huddle.
Never leave a ward with fewer than two registered nurses on the floor.

6. VISITORS

Visiting ends at 20:00 on every ward.
A relative may stay overnight with a patient who is dying or confused, with the nurse in charge's agreement.
Give them a fold-out chair and a visitor's pass from the drawer under the station desk.

7. FIRE

The fire assembly point for wards 3 to 9 is the staff car park by the incinerator.
Close every door you pass. Do not use the lifts.
Horizontal evacuation comes first: move patients through two sets of fire doors before going downstairs.
The night fire warden is the site manager unless the rota says otherwise.

8. IT SYSTEMS

If the electronic record is down, use the paper observation charts in the blue folder.
Enter the paper charts into the record within two hours of it coming back.
The IT service desk is staffed overnight only for problems that stop patient care.

9. SUPPLIES

The night store on the ground floor holds linen, fluids and dressings.
Its door code changes on the first day of each month and is in the site manager's handover.
Order blood gas cartridges before 02:00 so that they arrive with the morning delivery.

10. END OF SHIFT

Write up the night report before 07:15 so that the day team can read it at 07:30.
Return the controlled drug keys to the site manager in person.
"""

TEXT_ZH = """冷链仓库值班手册
2026年9月修订。本手册放在值班室，交班时一并移交。

一、值班安排

班次    负责人    交接时间    对讲机频道
早班    刘建国    08:00       3
中班    孙雅琴    16:00       5
夜班    马文涛    00:00       8

周末由仓储主管轮流带班，名单张贴在值班室门口。

二、交接清单

交班双方共同完成以下事项：
  - 核对冷库温度记录，冷冻区不得高于零下18度；
  - 检查叉车电量，低于30%时立即充电；
  - 清点当班出库单数量，并在系统中确认；
  - 检查月台门是否全部关闭并上锁。

清单未签字前，接班人员不得开始作业。

三、异常处理

冷冻区温度连续15分钟高于零下15度时，立即通知设备工程师。
停电超过10分钟时，启动备用发电机并记录启动时间。
发现货物外包装破损时，拍照后放入隔离区，不得出库。

四、门禁配置

门禁系统由安保部维护。如需重置，参数如下：

    DOOR_GROUP=coldchain-night
    ALARM_DELAY=90
    NOTIFY=security-desk
    AUTO_LOCK=22:00

门打开超过ALARM_DELAY秒仍未关闭时，系统会通知安保值班台。

五、来访车辆

供应商车辆须提前一天预约卸货时段，未预约车辆一律在门外等候。
冷藏车进入月台前须出示温度记录单，车厢温度高于零下12度的不予接收。
外来司机不得进入冷库作业区，卸货由本库叉车司机完成。

六、消防安全

冷库内严禁使用明火，维修动火须办理动火审批单。
消防通道每班巡查一次，堆放物品距离通道不少于1米。
灭火器每月15日由安保部检查并在标签上签字。

七、设备点检

制冷机组每两小时巡检一次，记录吸气压力和排气温度。
叉车充电区每天下班前清理，充电时不得离人。
月台升降平台每周一由维修班润滑保养。

八、交班

夜班须在早上7点30分前填写值班日志，并将冷库钥匙当面交给早班负责人。
"""


# ---------------------------------------------------------------------------
# Text PDFs: two columns, a table, footnotes, figures with captions
# ---------------------------------------------------------------------------


def pdf_styles(language):
    font = "STSong-Light" if language == "zh" else "Helvetica"
    bold = "STSong-Light" if language == "zh" else "Helvetica-Bold"
    italic = "STSong-Light" if language == "zh" else "Helvetica-Oblique"
    wrap = "CJK" if language == "zh" else None

    def style(name, **kw):
        return ParagraphStyle(name, wordWrap=wrap, **kw)

    return {
        "title": style("title", fontName=bold, fontSize=18, leading=23, spaceAfter=6),
        "byline": style("byline", fontName=font, fontSize=10, leading=13, spaceAfter=10),
        "abstract": style("abstract", fontName=font, fontSize=9.5, leading=13, spaceAfter=10),
        "h": style("h", fontName=bold, fontSize=11.5, leading=15, spaceBefore=6, spaceAfter=4),
        "p": style("p", fontName=font, fontSize=9.5, leading=13, spaceAfter=6),
        "caption": style("caption", fontName=italic, fontSize=8.5, leading=11, spaceBefore=3, spaceAfter=10),
        "cell": style("cell", fontName=font, fontSize=9, leading=12),
        "note": ("Helvetica" if language == "en" else "STSong-Light", 7.5),
    }


def bar_figure(categories, values, width, height):
    drawing = Drawing(width, height)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 28, 22, width - 40, height - 32
    chart.data = [values]
    chart.categoryAxis.categoryNames = categories
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 6.5
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 6.5
    chart.valueAxis.valueMin = 0
    chart.bars[0].fillColor = colors.HexColor("#7A8CA8")
    drawing.add(chart)
    return drawing


def map_figure(width, height, label):
    drawing = Drawing(width, height)
    drawing.add(Rect(0, 0, width, height, fillColor=colors.HexColor("#E8EEF2"), strokeColor=colors.grey))
    for x, y in [(0.2, 0.3), (0.35, 0.62), (0.5, 0.45), (0.62, 0.7), (0.74, 0.28), (0.82, 0.55)]:
        drawing.add(Rect(width * x, height * y, 5, 5, fillColor=colors.HexColor("#C0392B"), strokeWidth=0))
    drawing.add(String(6, height - 12, label, fontName="Helvetica", fontSize=7))
    return drawing


def write_pdf(path, language, spec):
    styles = pdf_styles(language)
    margin = 18 * mm
    gap = 7 * mm
    width, height = A4
    column = (width - 2 * margin - gap) / 2
    top_height = 62 * mm
    body_height = height - 2 * margin

    def frame(name, x, y, w, h):
        return Frame(x, y, w, h, id=name, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)

    bottom = margin + 20 * mm  # room for footnotes
    first = [
        frame("top", margin, height - margin - top_height, width - 2 * margin, top_height),
        frame("left1", margin, bottom, column, height - margin - top_height - bottom - 4 * mm),
        frame("right1", margin + column + gap, bottom, column, height - margin - top_height - bottom - 4 * mm),
    ]
    later = [
        frame("left", margin, bottom, column, body_height - 20 * mm),
        frame("right", margin + column + gap, bottom, column, body_height - 20 * mm),
    ]
    wide = [frame("wide", margin, bottom, width - 2 * margin, body_height - 20 * mm)]

    def decorate(canvas, document):
        canvas.saveState()
        font_name, size = styles["note"]
        page = document.page
        canvas.setFont(font_name, 7.5)
        canvas.drawRightString(width - margin, margin - 6 * mm, f"{spec['running']} · {page}")
        notes = spec["footnotes"].get(page, [])
        y = margin + 4 * mm + 3.6 * mm * len(notes)
        if notes:
            canvas.setLineWidth(0.4)
            canvas.line(margin, y + 4 * mm, margin + 45 * mm, y + 4 * mm)
        for number, text in notes:
            canvas.setFont(font_name, size)
            canvas.drawString(margin, y, f"{number} {text}")
            y -= 3.6 * mm
        canvas.restoreState()

    document = BaseDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        invariant=1,
        title="",
        author="",
        creator="",
        subject="",
    )
    document.addPageTemplates(
        [
            PageTemplate(id="first", frames=first, onPage=decorate),
            PageTemplate(id="columns", frames=later, onPage=decorate),
            PageTemplate(id="wide", frames=wide, onPage=decorate),
        ]
    )
    flow = [
        Paragraph(spec["title"], styles["title"]),
        Paragraph(spec["byline"], styles["byline"]),
        Paragraph(spec["abstract"], styles["abstract"]),
        NextPageTemplate("columns"),
    ]
    for kind, value in spec["columns"]:
        if kind == "h":
            flow.append(Paragraph(value, styles["h"]))
        elif kind == "p":
            flow.append(Paragraph(value, styles["p"]))
        elif kind == "figure":
            categories, values, caption = value
            flow.append(bar_figure(categories, values, column, 46 * mm))
            flow.append(Paragraph(caption, styles["caption"]))
    flow += [NextPageTemplate("wide"), PageBreak()]
    for kind, value in spec["wide"]:
        if kind == "h":
            flow.append(Paragraph(value, styles["h"]))
        elif kind == "p":
            flow.append(Paragraph(value, styles["p"]))
        elif kind == "table":
            caption, rows = value
            flow.append(Paragraph(caption, styles["caption"]))
            data = [[Paragraph(cell, styles["cell"]) for cell in row] for row in rows]
            table = Table(data, colWidths=[(width - 2 * margin) / len(rows[0])] * len(rows[0]), hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("LINEABOVE", (0, 0), (-1, 0), 0.8, colors.black),
                        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
                        ("LINEBELOW", (0, -1), (-1, -1), 0.8, colors.black),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            flow += [table, Spacer(1, 10)]
        elif kind == "map":
            label, caption = value
            flow.append(map_figure(width - 2 * margin, 52 * mm, label))
            flow.append(Paragraph(caption, styles["caption"]))
    document.build(flow)


PDF_EN = {
    "title": "Urban Heat in Six Districts: Summer 2026 Field Study",
    "byline": "Mara Lindqvist and Daniel Achterberg · City Climate Unit, Working Paper 14 (fictional)",
    "abstract": (
        "Abstract. We measured surface and air temperatures in six districts of a coastal city from June to "
        "August 2026 and tested four cooling measures. Districts with little tree cover were the hottest by day "
        "and cooled least at night. Misting stations gave the largest local cooling but cost the most to run."
    ),
    "running": "Urban Heat in Six Districts",
    "columns": [
        ("h", "1 Introduction"),
        ("p", "Heat waves are now the deadliest weather hazard in the city. Older residents in flats without "
              "cross-ventilation are most at risk, and ambulance call-outs rose by a fifth on the hottest days "
              "of 2025.<super>1</super>"),
        ("p", "The study asked two questions: which districts stay hottest, and which low-cost measures cool "
              "them most for each pound spent."),
        ("h", "2 Methods"),
        ("p", "Air temperature was logged every ten minutes in each district, and surface temperature was "
              "measured with a handheld infrared camera at noon on twenty clear days."),
        ("p", "District boundaries follow the census wards, merged where a ward had fewer than 4,000 "
              "residents.<super>2</super> Each measure was tested on at least two streets, with a "
              "matched street nearby as a control."),
        ("h", "3 Results"),
        ("p", "Surface temperatures in the Eastgate district peaked at 47.2 °C on 19 July, the highest of the "
              "six districts measured. Tree cover in Eastgate is 6 per cent, against a city average of "
              "19 per cent."),
        ("p", "Night-time air temperatures stayed above 24 °C for nine nights in a row in the Old Market area, "
              "where most streets run east to west and trap the evening sun."),
        ("figure", (
            ["Eastgate", "Old Market", "Docks", "Hillside", "Parkway", "Northfield"],
            [9.5, 8.0, 6.5, 4.0, 3.5, 3.0],
            "Figure 1. Hours above 30 °C on the hottest day, by district; Eastgate stayed above 30 °C for "
            "9.5 hours.",
        )),
        ("p", "The Hillside district, which has the most tree cover at 31 per cent, was on average 3.4 °C "
              "cooler at noon than Eastgate. The difference shrank to 1.1 °C after sunset."),
        ("p", "Wind from the sea cooled the Docks district by about 2 °C on most afternoons, but only within "
              "300 metres of the shore."),
        ("h", "4 Cooling measures"),
        ("p", "Cool roofs, painted with a white reflective coating, were the cheapest measure per square "
              "metre. Street trees took two summers to give shade but cooled the pavement beneath them the "
              "most."),
        ("p", "Misting stations cooled the air within five metres by up to 3.1 °C, but each used about 40 "
              "litres of water an hour.<super>3</super> Residents asked for them at bus stops and outside "
              "the two clinics."),
        ("figure", (
            ["Cool roofs", "Trees", "Misting", "Paving"],
            [1.8, 2.6, 3.1, 0.9],
            "Figure 2. Cooling beside each measure at noon, in °C, averaged over 14 test streets; the "
            "paving was measured in the Parkway car park.",
        )),
        ("p", "A survey of 1,140 residents found that 41 per cent of those in Eastgate changed their walking "
              "route on hot days, against 17 per cent citywide."),
        ("h", "5 Discussion"),
        ("p", "The hottest districts are also the poorest: median household income in Eastgate is 38 per cent "
              "below the city's. Heat adds to health risks that are already higher there, and fewer homes have "
              "fans or blinds."),
        ("p", "Tree planting is slow to pay back but cheap once trees are established. In our model a street "
              "tree costs about £22 a year to look after from its fifth year, mostly for watering in dry "
              "summers."),
        ("p", "Cool roofs suit the flat-roofed warehouses of the Docks district best. On pitched roofs the "
              "coating wore off within a single summer on two of the five test houses."),
        ("p", "Misting is a stopgap. It helps people waiting at bus stops on the hottest afternoons, but it "
              "raises humidity, and the stations need daily checks for legionella."),
        ("h", "6 Limitations"),
        ("p", "Twenty measurement days cover only part of the summer, and July 2026 was hotter than usual. "
              "Sensors in direct sun read up to 1.5 °C high until radiation shields were fitted on 3 July."),
        ("p", "The survey over-represents residents who were at home in the daytime, who are also those most "
              "exposed to indoor heat. Answers were collected on paper and online, in English and Polish."),
        ("p", "We did not measure indoor temperatures, which matter most for older residents at night; a "
              "follow-up study with 80 indoor loggers starts in June 2027."),
    ],
    "wide": [
        ("h", "7 Costs and recommendations"),
        ("table", (
            "Table 1. Cooling measures tested in summer 2026, with the area treated, the cooling measured and "
            "the cost.<super>4</super>",
            [
                ["Measure", "Area treated", "Cooling (°C)", "Cost"],
                ["Cool roofs", "18,400 m²", "1.8", "£24 per m²"],
                ["Street trees", "2,150 trees", "2.6", "£310 per tree"],
                ["Misting stations", "12 sites", "3.1", "£9,800 per site"],
                ["Permeable paving", "6,700 m²", "0.9", "£41 per m²"],
            ],
        )),
        ("p", "We recommend planting street trees in Eastgate and Old Market first, and keeping misting "
              "stations for bus stops and clinics, where people wait outside."),
        ("map", (
            "Misting stations, June 2026",
            "Figure 3. The 12 misting stations installed in June 2026, clustered around bus stops on the "
            "two main roads into Eastgate.",
        )),
        ("p", "The next study will repeat the measurements in summer 2027, after the first 600 trees are "
              "planted."),
    ],
    "footnotes": {
        1: [
            ("1", "Ambulance figures from the regional ambulance service's 2025 heat report."),
            ("2", "Wards were merged using the 2021 census boundaries."),
            ("3", "Water use was read from each station's meter weekly."),
        ],
        3: [("4", "Costs are 2026 prices and exclude maintenance, which adds about 8 per cent a year.")],
    },
}

PDF_ZH = {
    "title": "城市绿地与夏季降温调查（2026年）",
    "byline": "林雨桐、周启明 · 江城市生态环境研究所工作报告第9号（虚构）",
    "abstract": (
        "摘要：2026年6月至8月，我们在江城市五个城区测量了地表温度和气温，并评估了四种降温措施。"
        "绿地率低的城区白天最热、夜间降温最慢；屋顶绿化单位面积成本最高，但对室内降温效果最明显。"
    ),
    "running": "城市绿地与夏季降温调查",
    "columns": [
        ("h", "一、研究背景"),
        ("p", "近年来江城市夏季高温日数逐年增加，2025年日最高气温超过35℃的天数达到41天，"
              "比十年前多了12天。<super>1</super>"),
        ("p", "本研究关注两个问题：哪些城区高温最严重，以及哪些措施能以较低成本有效降温。"),
        ("h", "二、研究方法"),
        ("p", "在每个城区布设气温记录仪，每十分钟记录一次；地表温度在二十个晴天的正午用红外热像仪测量。"),
        ("p", "城区边界按街道行政区划确定，人口不足五千的街道与相邻街道合并。<super>2</super>"
              "每项措施至少在两条道路上试验，并在附近选择一条条件相近的道路作为对照。"),
        ("h", "三、调查结果"),
        ("p", "老城区正午地表温度最高达到46.8℃，出现在7月22日，是五个城区中最高的。"
              "老城区的绿地率仅为8%，而全市平均为23%。"),
        ("p", "临江区夜间气温连续七晚高于25℃，主要原因是高层建筑密集，夜间散热困难。"),
        ("figure", (
            ["Laocheng", "Linjiang", "Gaoxin", "Xihu", "Beishan"],
            [10.0, 8.5, 6.0, 4.5, 3.0],
            "图1 最热一天各城区气温高于30℃的小时数；老城区持续高于30℃达10小时。",
        )),
        ("p", "西湖区绿地率最高，达到38%，正午平均气温比老城区低3.8℃，日落后差距缩小到1.3℃。"),
        ("p", "江风使滨江一带下午气温平均降低约2℃，但影响范围只在离江岸四百米以内。"),
        ("h", "四、降温措施"),
        ("p", "屋顶绿化单位面积成本最高，但顶层住户室内气温平均下降2.2℃。"
              "行道树需要两个夏季才能形成树荫，但树下路面降温最明显。"),
        ("p", "喷雾降温设施可使周边五米内气温降低最多3.4℃，但每套每小时耗水约45升。<super>3</super>"
              "居民希望在公交站和菜市场门口设置。"),
        ("figure", (
            ["Green roofs", "Trees", "Misting", "Paving"],
            [2.2, 2.9, 3.4, 1.0],
            "图2 各项措施旁正午实测降温幅度（℃），为12条试验道路的平均值；透水铺装在高新区体育中心停车场测量。",
        )),
        ("p", "对1,260名居民的问卷显示，老城区有46%的受访者在高温天改变了步行路线，全市平均为19%。"),
        ("h", "五、讨论"),
        ("p", "高温最严重的城区往往也是老旧小区集中的地方。老城区1990年以前建成的住宅占62%，"
              "多数没有外墙保温，顶层住户夏季室内温度明显偏高。"),
        ("p", "行道树前期投入回收慢，但成活后养护成本低。按本研究测算，种植满五年后每棵树每年养护费用约为180元，"
              "主要用于旱季浇水和修剪。"),
        ("p", "屋顶绿化更适合高新区的平屋顶厂房；在老城区的坡屋顶住宅上，试点的六处中有两处出现渗漏，已停止推广。"),
        ("p", "喷雾降温只是临时措施，能缓解公交站候车乘客的炎热，但会增加空气湿度，设施需要每天检查水质。"),
        ("h", "六、研究局限"),
        ("p", "二十个测量日只覆盖了夏季的一部分，且2026年7月气温偏高。7月5日加装防辐射罩之前，"
              "阳光直射下的记录仪读数最多偏高1.6℃。"),
        ("p", "问卷样本中白天在家的居民比例偏高，而这部分居民恰恰受室内高温影响最大。"),
        ("p", "本研究没有测量室内温度。2027年6月将启动后续研究，在100户居民家中布设室内温度记录仪。"),
        ("h", "附：各城区测点说明"),
        ("p", "老城区：测点设在解放街、鼓楼巷和东门菜市场，三处均为东西走向的窄街，两侧以六层砖混住宅为主，"
              "沿街几乎没有行道树，午后整条街道都暴露在阳光下。"),
        ("p", "临江区：测点设在江景大道和滨江一路，周边为三十层以上的高层住宅，楼间距较小，"
              "夜间风速明显低于其他城区，傍晚以后散热缓慢。"),
        ("p", "高新区：测点设在科技园南门和体育中心，周边多为大面积硬化地面和平屋顶厂房，"
              "白天升温快，但夜间降温也快，昼夜温差是五个城区中最大的。"),
        ("p", "西湖区：测点设在植物园北路和茶园新村，周边绿地连片，并有水面调节，"
              "是本研究的对照城区，也是唯一夜间气温低于24℃的城区。"),
        ("p", "北山区：测点设在北山公园和矿机厂旧址，地势较高，夏季多偏北风，"
              "但测量期间有两处记录仪被移动过位置，相关数据已从分析中剔除。"),
        ("p", "所有记录仪安装高度为距地面2.5米，安装位置避开空调外机和汽车尾气直排口，"
              "每两周由研究助理现场核对一次时间和读数。"),
    ],
    "wide": [
        ("h", "七、成本与建议"),
        ("table", (
            "表1 2026年夏季试验的降温措施：实施规模、实测降温幅度和成本。<super>4</super>",
            [
                ["措施", "实施规模", "降温幅度（℃）", "成本"],
                ["屋顶绿化", "9,600平方米", "2.2", "每平方米380元"],
                ["行道树", "3,400棵", "2.9", "每棵2,600元"],
                ["喷雾降温设施", "15处", "3.4", "每处7.5万元"],
                ["透水铺装", "12,800平方米", "1.0", "每平方米260元"],
            ],
        )),
        ("p", "建议优先在老城区和临江区补种行道树，喷雾降温设施主要设在公交站和菜市场等人员等候的地方。"),
        ("map", (
            "Misting sites, June 2026",
            "图3 2026年6月建成的15处喷雾降温设施分布，集中在老城区两条主干道沿线的公交站。",
        )),
        ("p", "下一轮调查将在2027年夏季进行，届时首批800棵行道树已完成种植。"),
    ],
    "footnotes": {
        1: [
            ("1", "高温日数据来自江城市气象台2025年气候公报。"),
            ("2", "街道合并依据2023年行政区划调整方案。"),
            ("3", "耗水量每周从各设施水表读取一次。"),
        ],
        3: [("4", "成本按2026年价格计算，不含维护费用，维护费每年约为建设成本的6%。")],
    },
}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    write_docx(OUT / "Riverside Library Renovation Report.docx", "Riverside Library Renovation Report", WORD_EN)
    write_docx(OUT / "社区食堂试点评估报告.docx", "社区食堂试点评估报告", WORD_ZH)
    write_pptx(OUT / "Coffee Subscription Launch Review.pptx", DECK_EN)
    write_pptx(OUT / "新能源公交季度运营汇报.pptx", DECK_ZH)
    write_xlsx(OUT / "Clinic Staffing Plan 2027.xlsx", staffing_workbook())
    write_xlsx(OUT / "门店销售与库存2026.xlsx", sales_workbook())
    (OUT / "Bike Share Stations September 2026.csv").write_text(stations_csv(), encoding="utf-8")
    (OUT / "小区垃圾分类统计2026年9月.csv").write_text(sorting_csv(), encoding="utf-8")
    (OUT / "Field Kit Setup Guide.md").write_text(MARKDOWN_EN, encoding="utf-8")
    (OUT / "实验室数据管理规范.md").write_text(MARKDOWN_ZH, encoding="utf-8")
    (OUT / "Night Shift Handbook.txt").write_text(TEXT_EN, encoding="utf-8")
    (OUT / "冷链仓库值班手册.txt").write_text(TEXT_ZH, encoding="utf-8")
    write_pdf(OUT / "Urban Heat in Six Districts.pdf", "en", PDF_EN)
    write_pdf(OUT / "城市绿地与夏季降温调查.pdf", "zh", PDF_ZH)
    for path in sorted(OUT.iterdir()):
        print(f"{path.stat().st_size:>8}  {path.name}")


if __name__ == "__main__":
    main()
